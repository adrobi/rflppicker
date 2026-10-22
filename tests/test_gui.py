import os
import sys
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import openpyxl
from PyQt6.QtCore import QSettings
from PyQt6.QtWidgets import QApplication, QMessageBox

import qtprimer3_vis
from app_settings import AppSettings
from tests.helpers import HEADER, core_row, dispose, write_text
from theme import ThemeController

FOUR_SNPS = ('chrom\tpos\tref\talt\tid\n11\t37525128\tC\tG\tHOXB13\n1\t41250567\tA\tC\tLEPR\n'
             '1\t91308651\tC\tT\tOLFML3\n12\t880121\tA\tG\trs7781\n')
HOX = dict(variant='11:37525128 C>G', mapped_id='NC_056064.1', enzyme='Bse1I', isoschizomers='BsrI',
           genotype_quality='robust', genotype_margin=15.77, bands_ref_ref='448', bands_alt_alt='212; 236',
           product_size=475, assay_rank=1, primers_left='TATGAGCTTCTCTCCGCCAG', primers_right='CGCCTCCAAAGTAGCCGTAA',
           tm_left=59.0, tm_right=60.1)
FIXED_OUTPUT = (
    HEADER,
    [core_row(**HOX),
     core_row(**dict(HOX, enzyme='BsrDI', isoschizomers='', genotype_margin=12.1, assay_rank=2)),
     core_row(variant='1:41250567 A>C', enzyme='AciI', genotype_quality='nominal_only', genotype_margin=1.4,
              bands_ref_ref='292; 140', bands_alt_alt='432', product_size=432, assay_rank=1,
              weakest_genotypes='REF/REF–REF/ALT'),
     core_row(variant='1:91308651 C>T', enzyme='', status='no_discriminating_assay',
              reason='Не найдено сочетание для различения трёх генотипов')],
    {'summary': {'rows': 4}, 'rejected': [{'chrom': '12', 'pos': 880121, 'ref': 'A', 'alt': 'G',
                                            'reason': 'ref_not_match', 'details': 'x'}]},
)


class MainWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.tmp = self._dir.name
        self.ini = os.path.join(self.tmp, 's.ini')
        self.qs = QSettings(self.ini, QSettings.Format.IniFormat)
        self.settings = AppSettings(self.qs)
        self.fasta = write_text(self.tmp, 'genome.fa', '>1\nACGT\n')
        self.window = qtprimer3_vis.RFLPPickerGUI(self.settings, ThemeController(self.app, 'light'))
        self.asked = []

    def tearDown(self):
        self.window.details.close_fasta()
        dispose(self.window)
        self.qs.sync()
        self._dir.cleanup()

    def run_with(self, output, variants=FOUR_SNPS):
        path = write_text(self.tmp, 'variants.tsv', variants)
        self.window.launch.set_files(path, self.fasta)

        def fake(params, status_cb=None, progress_cb=None, cancel_cb=None, report=None):
            if isinstance(output, Exception):
                raise output
            header, rows, fixed_report = output
            if report is not None:
                report.update(fixed_report)
            if progress_cb:
                progress_cb(1, 4)
            return header, rows

        with patch('rflp_core.run_rflp_gui_mode', side_effect=fake):
            self.window.start_run()
            worker = self.window.worker
            self.assertIsNotNone(worker)
            worker.wait(10000)
            for _ in range(3):
                QApplication.processEvents()
        self.assertIsNone(self.window.worker)

    def test_run_fills_list_and_collapses_launch(self):
        self.run_with(FIXED_OUTPUT)
        self.assertEqual(self.window.list.model.rowCount(), 4)
        self.assertLessEqual(self.window.launch.maximumWidth(), 36)
        self.assertEqual(self.window.list.model.summary(0).record.snp_id, 'HOXB13')

    def test_selection_shows_details_and_choice_updates_the_list(self):
        self.run_with(FIXED_OUTPUT)
        self.window.list.table.selectRow(0)
        self.assertEqual(self.window.details.title_label.text(), 'HOXB13 · 11:37525128 C>G')
        self.window.details.choose_scheme(1)
        self.assertEqual(self.window.list.proxy.index(0, self.window.list.column_index('enzyme')).data(), 'BsrDI ✎')
        self.assertEqual(self.window.manual_choice_count(), 1)

    def test_rerun_asks_when_manual_choices_exist(self):
        self.run_with(FIXED_OUTPUT)
        self.window.list.model.summary(0).choose(1)
        self.window.confirm_rerun = lambda n: self.asked.append(n) or False
        self.window.start_run()
        self.assertEqual(self.asked, [1])
        self.assertIsNone(self.window.worker)

    def test_failed_run_keeps_previous_results(self):
        self.run_with(FIXED_OUTPUT)
        self.run_with(RuntimeError('boom'))
        self.assertEqual(self.window.list.model.rowCount(), 4)
        self.assertIn('boom', self.window.status_message.text())
        self.assertEqual(self.window.launch.run_button.text(), 'Запустить')

    def test_invalid_input_does_not_start(self):
        self.window.start_run()
        self.assertIsNone(self.window.worker)
        self.assertFalse(self.window.launch.findChild(object, 'error_input_file').isHidden())

    def test_export_shown_writes_only_filtered_snps(self):
        self.run_with(FIXED_OUTPUT)
        self.window.list.chips.toggled.emit({'robust'})
        path = self.window.export('shown', os.path.join(self.tmp, 'x.xlsx'))
        book = openpyxl.load_workbook(path)
        self.assertEqual(book['Сводка'].max_row, 2)
        self.assertEqual(book.sheetnames, ['Сводка', 'Праймеры', 'Results', 'Run', 'Rejected'])
        book.close()
        self.assertIn('Сохранено', self.window.status_message.text())

    def test_export_csv(self):
        self.run_with(FIXED_OUTPUT)
        path = self.window.export('csv', os.path.join(self.tmp, 'x.csv'))
        with open(path, encoding='utf-8-sig') as handle:
            self.assertTrue(handle.readline().rstrip().endswith('snp_id,selected'))

    def test_export_into_locked_file_warns(self):          # Review Focus 4
        self.run_with(FIXED_OUTPUT)
        target = write_text(self.tmp, 'x.xlsx', 'open in Excel')
        with patch('reporting.export_excel', side_effect=PermissionError()), \
                patch.object(QMessageBox, 'warning') as warn:
            self.assertIsNone(self.window.export('all', target))
        self.assertIn('Файл занят другой программой', warn.call_args[0][2])

    def test_export_into_a_write_protected_folder_says_so(self):
        self.run_with(FIXED_OUTPUT)
        with patch('reporting.export_excel', side_effect=PermissionError()), \
                patch.object(QMessageBox, 'warning') as warn:
            self.assertIsNone(self.window.export('all', os.path.join(self.tmp, 'new.xlsx')))
        self.assertIn('Нет доступа к папке', warn.call_args[0][2])

    def test_export_survives_an_uncreatable_results_folder(self):
        self.run_with(FIXED_OUTPUT)
        with patch.object(self.window, '_default_path', side_effect=PermissionError(13, 'denied')), \
                patch.object(QMessageBox, 'warning') as warn:
            self.assertIsNone(self.window.export('all'))
        self.assertIn('Нет доступа к папке', warn.call_args[0][2])

    def test_export_survives_unexpected_errors(self):
        self.run_with(FIXED_OUTPUT)
        with patch('reporting.export_excel', side_effect=ValueError('bad cell')), \
                patch.object(QMessageBox, 'warning') as warn:
            self.assertIsNone(self.window.export('all', os.path.join(self.tmp, 'x.xlsx')))
        self.assertIn('bad cell', warn.call_args[0][2])
        self.assertIn('ValueError', self.window.log.text.toPlainText())

    def test_exception_hook_logs_instead_of_aborting(self):
        previous = sys.excepthook
        self.addCleanup(setattr, sys, 'excepthook', previous)
        qtprimer3_vis.install_exception_hook(self.window)
        try:
            raise RuntimeError('slot failed')
        except RuntimeError as exc:
            with patch.object(QMessageBox, 'critical'):
                sys.excepthook(type(exc), exc, exc.__traceback__)
        self.assertIn('slot failed', self.window.status_message.text())
        self.assertIn('Traceback', self.window.log.text.toPlainText())

    def test_launch_panel_starts_expanded(self):
        self.settings.set_value('launch_collapsed', True)
        window = qtprimer3_vis.RFLPPickerGUI(self.settings, ThemeController(self.app, 'light'))
        self.addCleanup(dispose, window)
        self.assertFalse(window.launch.is_collapsed())

    def test_launch_stays_open_when_every_snp_is_an_error(self):
        rejected = [{'chrom': c, 'pos': p, 'ref': r, 'alt': a, 'reason': 'bad_chrom', 'details': 'x'}
                    for c, p, r, a in (('11', 37525128, 'C', 'G'), ('1', 41250567, 'A', 'C'),
                                       ('1', 91308651, 'C', 'T'), ('12', 880121, 'A', 'G'))]
        self.run_with((HEADER, [], {'rejected': rejected}))
        self.assertEqual(self.window.list.model.counts()['error'], 4)
        self.assertFalse(self.window.launch.is_collapsed())

    def test_theme_menu_switches_mode_and_is_saved(self):
        self.window.set_theme_mode('dark')
        self.assertEqual(self.window.theme_button.text(), 'Тема: тёмная ▾')
        self.window.close()
        self.assertEqual(AppSettings(QSettings(self.ini, QSettings.Format.IniFormat)).value('theme_mode'), 'dark')

    def test_empty_input_file(self):
        self.run_with(([], [], {}), variants='chrom\tpos\tref\talt\n')
        self.assertEqual(self.window.status_message.text(), 'В файле нет SNP')
        self.assertEqual(self.window.list.model.rowCount(), 0)

    def test_window_size_limits(self):
        self.assertEqual((self.window.minimumWidth(), self.window.minimumHeight()), (1000, 640))
        self.assertTrue(self.window.windowTitle().startswith('RFLP Picker'))

    def test_language_switch_is_saved_and_needs_restart(self):
        self.assertEqual(self.window.language_button.text(), 'Язык: системный ▾')
        self.assertEqual([a.text() for a in self.window.language_button.menu().actions()],
                         ['Системный', 'Русский', 'English'])
        with patch.object(QMessageBox, 'information') as info:
            self.window.set_language_mode('en')
        self.assertEqual(self.settings.value('language'), 'en')
        self.assertIn('перезапуска', info.call_args[0][2])
        self.assertEqual(self.window.language_button.text(), 'Язык: English ▾')
        with patch.object(QMessageBox, 'information') as info:
            self.window.set_language_mode('ru')          # the language of this session: nothing to restart
        info.assert_not_called()
        self.assertEqual(self.settings.value('language'), 'ru')

    def test_startup_language(self):                     # Review Focus 2
        self.settings.set_value('language', 'fr')
        with patch('qtprimer3_vis.QLocale') as locale:
            locale.system.return_value.uiLanguages.return_value = ['de-DE']
            self.assertEqual(qtprimer3_vis.startup_language(self.settings), 'en')
            locale.system.return_value.uiLanguages.return_value = ['ru-RU', 'en-US']
            self.assertEqual(qtprimer3_vis.startup_language(self.settings), 'ru')
            self.settings.set_value('language', 'en')
            self.assertEqual(qtprimer3_vis.startup_language(self.settings), 'en')
            self.settings.set_value('language', 'system')
            self.assertEqual(qtprimer3_vis.startup_language(self.settings), 'ru')


if __name__ == '__main__':
    unittest.main()
