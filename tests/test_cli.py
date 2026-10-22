import contextlib
import csv
import io
import logging
import random
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import i18n
import rflp_cli
import rflp_core as core


class CliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.addCleanup(i18n.set_language, 'ru')
        rng = random.Random(3)
        sequence = ''.join(rng.choice('ACGT') for _ in range(400))
        sequence = sequence[:197] + 'GAATTC' + sequence[203:]
        self.fasta = self.root / 'toy.fa'
        self.fasta.write_text('>1\n' + sequence + '\n')
        self.variants = self.root / 'variants.tsv'
        self.variants.write_text('chrom\tpos\tref\talt\tid\n1\t200\tA\tG\tSNP1\n1\t2\tA\tA\tSNP2\n')

    def run_cli(self, *args):
        # Russian unless a test asks otherwise: the language of the CI machine is English.
        stderr = io.StringIO()
        with contextlib.redirect_stderr(stderr):
            code = rflp_cli.main(['--lang', 'ru', '--fasta', str(self.fasta), '--variants', str(self.variants),
                                  *args])
        return code, stderr.getvalue()

    def test_csv_matches_core_run(self):
        out = self.root / 'out.csv'
        rejected = self.root / 'rejected.csv'
        code, log = self.run_cli('--out', str(out), '--rejected', str(rejected), '--no-primer3',
                                 '--flank', '150', '--min-frag', '20', '--preset', 'v1.1')
        self.assertEqual(code, 0, log)
        params = dict(core.DEFAULT_PARAMS, fasta=str(self.fasta.resolve()),
                      input_file=str(self.variants.resolve()), use_primer3=False, flank=150, min_frag=20)
        header, rows = core.run_rflp_gui_mode(params)
        with out.open(encoding='utf-8-sig', newline='') as stream:
            written = list(csv.reader(stream))
        self.assertEqual(written[0], header + ['snp_id', 'selected'])
        self.assertEqual([line[:len(header)] for line in written[1:]],
                         [['' if v is None else str(v) for v in row] for row in rows])
        self.assertEqual({line[-2] for line in written[1:]}, {'SNP1'})
        self.assertTrue(any(r['enzyme'] == 'EcoRI' for r in (dict(zip(header, row)) for row in rows)))
        with rejected.open(encoding='utf-8-sig', newline='') as stream:
            self.assertEqual(len(list(csv.DictReader(stream))), 1)
        self.assertIn('Сохранено', log)

    def test_excel_has_summary_primer_and_technical_sheets(self):
        out = self.root / 'out.xlsx'
        code, log = self.run_cli('--out', str(out), '--no-primer3', '--min-frag', '20', '-q')
        self.assertEqual((code, log), (0, ''))
        workbook = load_workbook(out, read_only=True)
        self.assertEqual(workbook.sheetnames, ['Сводка', 'Праймеры', 'Results', 'Run', 'Rejected'])
        summary = list(workbook['Сводка'].iter_rows(values_only=True))
        self.assertEqual([row[0] for row in summary], ['ID', 'SNP1', 'SNP2'])
        self.assertEqual(summary[2][3], 'Ошибка данных')
        results_header = next(workbook['Results'].iter_rows(max_row=1, values_only=True))
        self.assertEqual(list(results_header[-2:]), ['snp_id', 'selected'])
        run = {row[0]: row[1] for row in workbook['Run'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(run['app_version'], core.APP_VERSION)
        workbook.close()

    def test_invalid_parameters_and_extension(self):
        code, log = self.run_cli('--out', str(self.root / 'x.csv'), '--prod-min', '700', '--prod-max', '600')
        self.assertEqual(code, 2)
        self.assertIn('размер продукта', log)
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            self.run_cli('--out', str(self.root / 'x.txt'))

    def test_english_help_has_no_cyrillic(self):
        out = io.StringIO()
        with contextlib.redirect_stdout(out), self.assertRaises(SystemExit):
            rflp_cli.main(['--lang', 'en', '--help'])
        self.assertIsNone(re.search('[А-Яа-яЁё]', out.getvalue()), out.getvalue())
        self.assertIn('--lang', out.getvalue())

    def test_english_excel_and_messages(self):
        out = self.root / 'out.xlsx'
        code, log = self.run_cli('--lang', 'en', '--out', str(out), '--no-primer3', '--min-frag', '20')
        self.assertEqual(code, 0, log)
        self.assertIn('Saved', log)
        self.assertIsNone(re.search('[А-Яа-яЁё]', log), log)
        book = load_workbook(out, read_only=True)
        self.assertEqual(book.sheetnames, ['Summary', 'Primers', 'Results', 'Run', 'Rejected'])
        self.assertEqual(next(book['Summary'].values)[:4], ('ID', 'SNP', 'Genome chromosome', 'Status'))
        book.close()

    def test_english_parameter_errors(self):
        code, log = self.run_cli('--lang', 'en', '--out', str(self.root / 'x.csv'), '--prod-min', '900',
                                 '--prod-max', '100')
        self.assertEqual(code, 2)
        self.assertEqual(log.strip(), 'Parameter error: Minimum product size is greater than maximum')

    def test_english_core_log_lines(self):
        records = []
        handler = logging.Handler()
        handler.emit = lambda record: records.append(record.getMessage())
        logger = logging.getLogger('RFLPPickerLogger')
        logger.addHandler(handler)
        self.addCleanup(logger.removeHandler, handler)
        self.variants.write_text('chrom\tpos\tref\talt\tid\n1\t200\tA\tG\tSNP1\n1\tx\tA\tG\tSNP3\n')
        code, log = self.run_cli('--lang', 'en', '--out', str(self.root / 'x.csv'), '--no-primer3',
                                 '--names', str(self.root / 'missing.txt'))
        self.assertEqual(code, 0, log)
        self.assertTrue(any(x.startswith('Names file not found') for x in records), records)
        self.assertTrue(any(x.startswith('Line parse error') for x in records), records)
        self.assertEqual([x for x in records if re.search('[А-Яа-яЁё]', x)], [])

    def test_unreadable_system_locale_falls_back_to_english(self):
        with patch('rflp_cli.locale.getlocale', side_effect=ValueError('unknown locale')), \
                contextlib.redirect_stdout(io.StringIO()), self.assertRaises(SystemExit) as stop:
            rflp_cli.main(['--help'])
        self.assertEqual(stop.exception.code, 0)
        self.assertEqual(i18n.language(), 'en')

    def test_lang_without_a_value_gets_the_full_usage(self):
        stderr = io.StringIO()
        with contextlib.redirect_stderr(stderr), self.assertRaises(SystemExit) as stop:
            rflp_cli.main(['--fasta', 'a', '--lang'])
        self.assertEqual(stop.exception.code, 2)
        self.assertIn('--variants VARIANTS', stderr.getvalue())

    def test_language_follows_the_system_without_the_key(self):
        with patch('rflp_cli.locale.getlocale', return_value=('en_US', 'UTF-8')), \
                contextlib.redirect_stdout(io.StringIO()), self.assertRaises(SystemExit):
            rflp_cli.main(['--help'])
        self.assertEqual(i18n.language(), 'en')
        with patch('rflp_cli.locale.getlocale', return_value=('Russian_Russia', '1251')), \
                contextlib.redirect_stdout(io.StringIO()), self.assertRaises(SystemExit):
            rflp_cli.main(['--help'])
        self.assertEqual(i18n.language(), 'ru')


if __name__ == '__main__':
    unittest.main()
