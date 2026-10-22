"""The whole interface in English: no Russian text left in the window or the Excel summary (i18n spec §1, §7.3–7.4)."""
import os
import tempfile
import unittest

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import openpyxl
from PyQt6.QtCore import QSettings
from PyQt6.QtWidgets import QApplication

import i18n
import qtprimer3_vis
from app_settings import AppSettings
from tests.helpers import CYRILLIC, FIXED_OUTPUT, cyrillic_texts, dispose, run_fake_analysis, write_text
from theme import ThemeController

CORE_LOG = [
    '\n[INFO] NORMALIZE: загружено вариантов: 4',
    '[INFO] Поиск: joint; SNP: 3; ферментов: 623',
    '[INFO] Оценка полос: модель присутствия, полная рестрикция; '
    'яркость и специфичность ПЦР по всему геному не оцениваются.',
    '[INFO] Отбор: только продаваемые ферменты; изошизомеры объединены (одна строка на сайт); '
    'dCAPS если нет robust-схемы, несовпадений до 1',
    '[INFO] 11:37525128 C>G: пар 10, кандидатов 2, неразличимых 0',
    '[WARN] В окнах SNP нет строчных букв: FASTA без soft-mask, маскирование повторов не действует.',
]



def russian_texts(widget) -> list:
    # The language menu names Russian in Russian (spec §3.2).
    return [text for text in cyrillic_texts(widget) if text != qtprimer3_vis.LANGUAGE_NAMES['ru']]


class EnglishInterfaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        i18n.set_language('en')
        self.addCleanup(i18n.set_language, 'ru')
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)
        self.tmp = self._dir.name
        self.qs = QSettings(os.path.join(self.tmp, 's.ini'), QSettings.Format.IniFormat)
        self.fasta = write_text(self.tmp, 'genome.fa', '>1\nACGT\n')
        self.window = qtprimer3_vis.RFLPPickerGUI(AppSettings(self.qs), ThemeController(self.app, 'light'))
        self.addCleanup(dispose, self.window)
        self.addCleanup(self.window.details.close_fasta)

    def test_window_has_no_russian_text(self):
        self.assertEqual(russian_texts(self.window), [])
        run_fake_analysis(self.window, self.tmp, self.fasta, FIXED_OUTPUT, log=CORE_LOG)
        self.assertTrue(self.window.status_message.text().startswith('Done: 4 SNPs'))
        for row in range(4):                     # robust, good gel only, no assay, data error
            self.window.list.table.selectRow(row)
            self.assertEqual(russian_texts(self.window), [], f'row {row}')
        self.window.list.table.selectRow(0)
        self.window.details.view_scheme(1)
        self.window.details.open_large()
        self.assertEqual(russian_texts(self.window.details.large_dialog()), [])
        self.assertIn('[INFO] Search: joint; SNPs: 3; enzymes: 623', self.window.log.text.toPlainText())

    def test_failed_run_message_is_english(self):
        run_fake_analysis(self.window, self.tmp, self.fasta, ValueError('Некорректное разрешение геля'))
        self.assertEqual(self.window.status_message.text(), 'Run failed: Invalid gel resolution')

    def test_english_excel_keeps_status_fills(self):            # Review Focus 3
        run_fake_analysis(self.window, self.tmp, self.fasta, FIXED_OUTPUT)
        path = self.window.export('all', os.path.join(self.tmp, 'x.xlsx'))
        book = openpyxl.load_workbook(path)
        self.assertEqual(book.sheetnames, ['Summary', 'Primers', 'Results', 'Run', 'Rejected'])
        summary = book['Summary']
        header = [cell.value for cell in summary[1]]
        status = header.index('Status') + 1
        for sheet in (summary, book['Primers']):
            texts = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
            self.assertEqual([x for x in texts if CYRILLIC.search(x)], [], sheet.title)
        fills = {summary.cell(row, status).value: summary.cell(row, status).fill.fgColor.rgb[-6:]
                 for row in range(2, summary.max_row + 1)}
        self.assertEqual(fills, {'Robust': 'EAF3DE', 'Good gel only': 'FAEEDA', 'No assay': 'EDEDED',
                                 'Data error': 'FCEBEB'})
        book.close()


if __name__ == '__main__':
    unittest.main()
