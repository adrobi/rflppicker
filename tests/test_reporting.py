"""Round-trip checks for typed, safe, complete exports."""

import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from reporting import collect_versions, export_csv, export_excel


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)
        self.addCleanup(self.temporary.cleanup)

    def workbook(self, filename="results.xlsx"):
        workbook = load_workbook(self.folder / filename, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def test_excel_preserves_numbers_unicode_identifiers_and_formula_text(self):
        header = ["variant", "mapped_id", "pos", "tm_left", "primers_left", "note"]
        rows = [["Овца:25:A:G", "001234567890123456789", 123456789, 61.375, "AGCTAGCTAGCT", "=1+1"]]
        export_excel(self.folder / "results.xlsx", header, rows)
        workbook = self.workbook()
        sheet = workbook["Results"]
        self.assertEqual([cell.value for cell in sheet[2]], rows[0])
        self.assertEqual(sheet["C2"].data_type, "n")
        self.assertEqual(sheet["C2"].number_format, "0")
        self.assertEqual(sheet["D2"].data_type, "n")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["B2"].number_format, "@")
        self.assertEqual(sheet["F2"].data_type, "s")
        self.assertEqual(sheet["E2"].font.name, "Consolas")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:F2")

    def test_excel_metadata_and_rejected_rows(self):
        metadata = {
            "genome_path": "D:/геномы/овца.fa",
            "run_timestamp": "2026-09-12T14:00:00+03:00",
            "parameters": {"flank": 500, "supplier": "Все"},
            "status_summary": {"success": 1, "rejected": 2},
            "comment": "=HYPERLINK(\"https://example.org\")",
        }
        rejected = [
            {"chrom": "1", "pos": 12, "reason": "Несовпадение REF"},
            {"chrom": "2", "pos": 18, "reason": "+проверить", "line": 3},
        ]
        export_excel(self.folder / "results.xlsx", ["pos"], [[25]], metadata, rejected)
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ["Results", "Run", "Rejected"])
        run = {row[0].value: row[1] for row in list(workbook["Run"].rows)[1:]}
        self.assertEqual(run["genome_path"].value, metadata["genome_path"])
        self.assertEqual(json.loads(run["parameters"].value), metadata["parameters"])
        self.assertEqual(run["result_count"].value, 1)
        self.assertEqual(run["rejected_count"].value, 2)
        self.assertEqual(run["comment"].data_type, "s")
        self.assertIn("Python", json.loads(run["versions"].value))
        sheet = workbook["Rejected"]
        self.assertEqual([cell.value for cell in sheet[1]], ["chrom", "pos", "reason", "line"])
        self.assertEqual(sheet["B2"].value, 12)
        self.assertEqual(sheet["C3"].value, "+проверить")
        self.assertEqual(sheet["C3"].data_type, "s")

    def test_empty_results_and_rejections_keep_headers(self):
        export_excel(self.folder / "results.xlsx", ["variant", "pos"], [])
        workbook = self.workbook()
        self.assertEqual(workbook["Results"].max_row, 1)
        self.assertEqual([cell.value for cell in workbook["Results"][1]], ["variant", "pos"])
        self.assertEqual(
            [cell.value for cell in workbook["Rejected"][1]], ["chrom", "pos", "ref", "alt", "reason"]
        )
        export_csv(self.folder / "results.csv", ["variant", "pos"], [])
        with (self.folder / "results.csv").open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(list(csv.reader(stream)), [["variant", "pos"]])

    def test_csv_unicode_all_rows_and_formula_protection(self):
        values = ["=1+1", "+SUM(1,2)", "-3", "@command", "  =2+2", "\tfoo", "\nfoo", "Овца", "AGCT"]
        rows = ([index, value, -5.25] for index, value in enumerate(values))
        path = self.folder / "results.csv"
        export_csv(path, ["pos", "text", "numeric"], rows)
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
        with path.open(encoding="utf-8-sig", newline="") as stream:
            exported = list(csv.reader(stream))
        self.assertEqual(len(exported), len(values) + 1)
        for index, value in enumerate(values):
            expected = "'" + value if index < 7 else value
            self.assertEqual(exported[index + 1], [str(index), expected, "-5.25"])

    def test_failed_csv_generator_preserves_destination(self):
        path = self.folder / "results.csv"
        path.write_text("previous result", encoding="utf-8")

        def rows():
            yield ["first"]
            raise RuntimeError("interrupted result source")

        with self.assertRaisesRegex(RuntimeError, "interrupted"):
            export_csv(path, ["variant"], rows())
        self.assertEqual(path.read_text(encoding="utf-8"), "previous result")
        self.assertEqual(list(self.folder.iterdir()), [path])

    def test_invalid_width_and_excel_truncation_do_not_overwrite(self):
        for suffix, exporter in (("csv", export_csv), ("xlsx", export_excel)):
            with self.subTest(suffix=suffix):
                path = self.folder / f"invalid.{suffix}"
                path.write_bytes(b"previous")
                with self.assertRaisesRegex(ValueError, "expected 2"):
                    exporter(path, ["a", "b"], [[1]])
                self.assertEqual(path.read_bytes(), b"previous")
        path = self.folder / "long.xlsx"
        path.write_bytes(b"previous")
        with self.assertRaisesRegex(ValueError, "cell limit"):
            export_excel(path, ["sequence"], [["A" * 32768]])
        self.assertEqual(path.read_bytes(), b"previous")

    def test_failed_excel_save_preserves_destination_and_removes_temporary_file(self):
        path = self.folder / "results.xlsx"
        path.write_bytes(b"previous workbook")

        def failed_save(_workbook, temporary_path):
            Path(temporary_path).write_bytes(b"incomplete ZIP data")
            raise OSError("disk full")

        with patch("openpyxl.workbook.workbook.Workbook.save", new=failed_save):
            with self.assertRaisesRegex(OSError, "disk full"):
                export_excel(path, ["pos"], [[125]])
        self.assertEqual(path.read_bytes(), b"previous workbook")
        self.assertEqual(list(self.folder.iterdir()), [path])

    def test_versions_include_runtime_and_installed_packages(self):
        versions = collect_versions()
        self.assertTrue(versions["Python"])
        self.assertNotEqual(versions["openpyxl"], "not installed")


if __name__ == "__main__":
    unittest.main()
