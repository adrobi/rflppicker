"""Safe, atomic CSV and Excel exports for RFLP Picker.

CSV is written as UTF-8 with a BOM so that Excel recognises Cyrillic text.
Because CSV carries no cell types, potentially executable text (``= + - @``
after whitespace, or text beginning with a tab/newline) is prefixed with an
apostrophe. This deliberately changes those CSV strings; numeric values are
unchanged. XLSX preserves the original text and explicitly marks it as text.

Both writers replace the destination only after a successful complete write.
Malformed row widths and Excel limits raise errors instead of losing data.
"""

from __future__ import annotations

import csv
import json
import math
import os
import platform
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from contextlib import contextmanager
from datetime import date, datetime, time, timezone
from decimal import Decimal
from importlib import metadata as package_metadata
from pathlib import Path
from typing import Any, Iterator


_REJECTED_HEADERS = ["chrom", "pos", "ref", "alt", "reason"]
_EXCEL_MAX_ROWS = 1_048_576
_EXCEL_MAX_COLUMNS = 16_384
_EXCEL_MAX_STRING = 32_767


def collect_versions() -> dict[str, str]:
    """Return the interpreter and relevant installed distribution versions."""
    versions = {"Python": platform.python_version()}
    for name in ("PyQt6", "pyfaidx", "biopython", "primer3-py", "openpyxl"):
        try:
            versions[name] = package_metadata.version(name)
        except package_metadata.PackageNotFoundError:
            versions[name] = "not installed"
    return versions


@contextmanager
def _atomic_destination(path: str | os.PathLike[str]) -> Iterator[Path]:
    destination = Path(path).expanduser().absolute()
    # The parent must already exist; a misspelled path should not create folders.
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=destination.suffix, dir=destination.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        yield temporary
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _headers(header: Sequence[Any]) -> list[str]:
    if isinstance(header, (str, bytes)):
        raise TypeError("header must be a sequence of column names")
    return [str(value) for value in header]


def _checked_rows(rows: Iterable[Sequence[Any]], width: int) -> Iterator[list[Any]]:
    for index, row in enumerate(rows, start=1):
        if isinstance(row, (str, bytes, Mapping)):
            raise TypeError(f"Row {index} must be a sequence of cell values")
        values = list(row)
        if len(values) != width:
            raise ValueError(f"Row {index} has {len(values)} values; expected {width}")
        yield values


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def _csv_value(value: Any) -> Any:
    if isinstance(value, (Mapping, list, tuple)):
        value = _json(value)
    if isinstance(value, str):
        stripped = value.lstrip()
        if value.startswith(("\t", "\r", "\n")) or stripped.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def export_csv(
    path: str | os.PathLike[str], header: Sequence[Any], rows: Iterable[Sequence[Any]]
) -> None:
    """Export every supplied row; dangerous text receives an apostrophe prefix.

    ``rows`` may be a generator. A failed generator/write leaves an existing
    destination intact. Column counts must match ``header`` exactly.
    """
    names = _headers(header)
    with _atomic_destination(path) as temporary:
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([_csv_value(value) for value in names])
            for row in _checked_rows(rows, len(names)):
                writer.writerow([_csv_value(value) for value in row])


def _excel_value(value: Any) -> Any:
    if isinstance(value, (Mapping, list, tuple)):
        value = _json(value)
    elif isinstance(value, (datetime, time)) and value.tzinfo is not None:
        value = value.isoformat()
    elif isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Excel cannot preserve non-finite numeric values")
    elif isinstance(value, Decimal) and not value.is_finite():
        raise ValueError("Excel cannot preserve non-finite numeric values")
    elif value is not None and not isinstance(value, (str, int, float, bool, Decimal, date, time)):
        value = str(value)
    if isinstance(value, str) and len(value) > _EXCEL_MAX_STRING:
        raise ValueError(f"Excel text exceeds the {_EXCEL_MAX_STRING}-character cell limit")
    return value


def _write_cell(worksheet: Any, row: int, column: int, value: Any) -> Any:
    value = _excel_value(value)
    cell = worksheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        # openpyxl treats strings beginning with '=' as formulas by default.
        cell.data_type = "s"
        cell.number_format = "@"
    elif isinstance(value, int) and not isinstance(value, bool):
        cell.number_format = "0"
    elif isinstance(value, (float, Decimal)):
        cell.number_format = "0.00############"
    return cell


def _populate_sheet(worksheet: Any, header: Sequence[str], rows: list[list[Any]], table_name: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    if len(header) > _EXCEL_MAX_COLUMNS or len(rows) + 1 > _EXCEL_MAX_ROWS:
        raise ValueError("The export exceeds Excel's maximum worksheet dimensions")
    if not header and rows:
        raise ValueError("Nonempty Excel data requires at least one column")
    if any(not name.strip() for name in header) or len({name.casefold() for name in header}) != len(header):
        raise ValueError("Excel column names must be nonempty and unique")
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:1"
    worksheet.row_dimensions[1].height = 30
    for column, name in enumerate(header, start=1):
        cell = _write_cell(worksheet, 1, column, name)
        cell.fill = PatternFill("solid", fgColor="183B56")
        cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_index, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = _write_cell(worksheet, row_index, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            column_name = header[column - 1].lower()
            sequence_column = "primer" in column_name and isinstance(value, str)
            cell.font = Font(name="Consolas" if sequence_column else "Calibri", size=11, color="23374D")
    for column, name in enumerate(header, start=1):
        lengths = [len(name)]
        lengths.extend(
            max((len(line) for line in str(values[column - 1] or "").splitlines()), default=0)
            for values in rows
        )
        width = min(64, max(13, max(lengths) + 2))
        if worksheet.title == "Run" and column == 2:
            width = 85
        worksheet.column_dimensions[get_column_letter(column)].width = width
    if header:
        reference = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
        worksheet.auto_filter.ref = reference
        if rows:
            table = Table(displayName=table_name, ref=reference)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            worksheet.add_table(table)


def _rejected_data(rejected: Iterable[Any] | None) -> tuple[list[str], list[list[Any]]]:
    items = list(rejected) if rejected is not None else []
    if not items:
        return list(_REJECTED_HEADERS), []
    if isinstance(items[0], Mapping):
        names: list[str] = []
        for item in items:
            if not isinstance(item, Mapping):
                raise TypeError("Rejected rows must all be mappings or all be sequences")
            for key in item:
                if not isinstance(key, str):
                    raise TypeError("Rejected mapping keys must be strings")
                if key not in names:
                    names.append(key)
        names = names or list(_REJECTED_HEADERS)
        return names, [[item.get(name) for name in names] for item in items]
    return list(_REJECTED_HEADERS), list(_checked_rows(items, len(_REJECTED_HEADERS)))


def export_excel(
    path: str | os.PathLike[str],
    header: Sequence[Any],
    rows: Iterable[Sequence[Any]],
    metadata: Mapping[str, Any] | None = None,
    rejected: Iterable[Mapping[str, Any] | Sequence[Any]] | None = None,
) -> None:
    """Write typed Results, Run, and Rejected worksheets atomically.

    Numeric inputs stay numeric; identifier strings and primer strings stay
    text, including strings that look like formulas. Pass typed model values,
    not preformatted table strings, to retain numeric cell types.

    ``metadata`` is any mapping (nested values are JSON). It should include
    analysis parameters, genome path, run timestamp and status summary.
    Export timestamp, row counts and dependency versions are added when absent.
    ``rejected`` accepts dictionaries with arbitrary string keys, or sequences
    in the order chrom, pos, ref, alt, reason.
    """
    from openpyxl import Workbook

    names = _headers(header)
    result_rows = list(_checked_rows(rows, len(names)))
    rejected_header, rejected_rows = _rejected_data(rejected)
    run_metadata = dict(metadata) if metadata is not None else {}
    run_metadata.setdefault("exported_at", datetime.now(timezone.utc).isoformat(timespec="seconds"))
    run_metadata.setdefault("result_count", len(result_rows))
    run_metadata.setdefault("rejected_count", len(rejected_rows))
    run_metadata.setdefault("versions", collect_versions())
    run_rows = [[str(key), value] for key, value in run_metadata.items()]

    workbook = Workbook()
    try:
        results_sheet = workbook.active
        results_sheet.title = "Results"
        _populate_sheet(results_sheet, names, result_rows, "RFLPResults")
        _populate_sheet(workbook.create_sheet("Run"), ["Parameter", "Value"], run_rows, "RFLPRun")
        _populate_sheet(workbook.create_sheet("Rejected"), rejected_header, rejected_rows, "RFLPRejected")
        with _atomic_destination(path) as temporary:
            workbook.save(temporary)
    finally:
        workbook.close()
